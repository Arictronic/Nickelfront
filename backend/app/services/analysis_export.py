import io
import json
import logging

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models.analysis_result import AnalysisResult

logger = logging.getLogger(__name__)


async def export_analysis_to_excel(
    analysis_id: int,
    db: AsyncSession,
    user_id: int | None = None,
) -> io.BytesIO:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required for Excel export. Install dependencies from requirements.txt."
        ) from exc

    stmt = select(AnalysisResult).where(AnalysisResult.id == analysis_id)
    if user_id is not None:
        stmt = stmt.where(AnalysisResult.user_id == user_id)

    result = await db.execute(stmt)
    entry = result.scalar_one_or_none()
    if not entry:
        raise ValueError(f"AnalysisResult {analysis_id} не найден")

    wb = Workbook()
    ws = wb.active
    ws.title = "Анализ"

    header_font = Font(bold=True, size=12)
    header_alignment = Alignment(horizontal="center")

    headers = ["Поле", "Значение"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = header_alignment

    fields = [
        ("ID анализа", entry.id),
        ("ID статьи", entry.paper_id),
        ("Статус", entry.status),
        ("Версия промпта", entry.prompt_version),
        ("Сообщение об ошибке", entry.error_message),
        ("Сырой ответ Qwen", entry.raw_response),
        ("Структурированный результат", str(entry.structured_result) if entry.structured_result else None),
        ("Создан", str(entry.created_at) if entry.created_at else None),
        ("Завершён", str(entry.completed_at) if entry.completed_at else None),
    ]

    for row_idx, (field_name, value) in enumerate(fields, 2):
        ws.cell(row=row_idx, column=1, value=field_name)
        ws.cell(row=row_idx, column=2, value=str(value) if value is not None else "-")

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 80

    if entry.structured_result and isinstance(entry.structured_result, dict):
        start_row = len(fields) + 3
        ws.cell(row=start_row, column=1, value="Структурированные данные").font = header_font
        start_row += 1
        for key, value in entry.structured_result.items():
            ws.cell(row=start_row, column=1, value=str(key))
            if isinstance(value, list):
                ws.cell(row=start_row, column=2, value="\n".join(str(v) for v in value))
            elif isinstance(value, dict):
                ws.cell(row=start_row, column=2, value=json.dumps(value, ensure_ascii=False, indent=2))
            else:
                ws.cell(row=start_row, column=2, value=str(value))
            start_row += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
